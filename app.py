import streamlit as st
from docx import Document
from docx.shared import RGBColor, Pt
import json
import io
import datetime

# Configuration de la page
st.set_page_config(page_title="Assistant PV Universel", layout="wide")

st.title("⚡ Assistant PV de Séance - Multi-Projets")

# --- 1. LE CERVEAU : IMPORT DU JSON ET PARAMÈTRES GÉNÉRAUX ---
with st.sidebar:
    st.header("📂 1. Configuration & Historique")
    
    nom_projet = st.text_input("Nom du Projet", "Projet Alpha")
    date_seance = st.date_input("Date de la séance", datetime.date.today())
    
    st.info("Glisse ici le fichier JSON de la séance précédente pour ce projet.")
    fichier_json = st.file_uploader("Importer l'historique", type=["json"])
    
    if fichier_json is not None and 'import_fait' not in st.session_state:
        historique = json.load(fichier_json)
        for p in historique:
            p['color'] = "#000000" # Force le noir pour les anciens points
        st.session_state.points = historique
        st.session_state.import_fait = True
        st.success("Ancien PV chargé. Les nouveautés sont passées en noir.")
        st.rerun()

# --- INITIALISATION DE BASE ---
if 'points' not in st.session_state:
    st.session_state.points = []

# --- 2. L'ACTION : AJOUT EN SÉANCE ---
st.subheader("➕ 2. Ajouter un point (Séance en cours)")

intervenants_existants = list(set([p["qui"] for p in st.session_state.points]))
if not intervenants_existants:
    intervenants_existants = ["Nouveau..."]
else:
    intervenants_existants = ["Nouveau..."] + sorted(intervenants_existants)

with st.form("ajout"):
    col1, col2, col3 = st.columns([1, 2, 1])
    
    qui_existant = col1.selectbox("Responsable", intervenants_existants)
    qui_nouveau = col1.text_input("Ou nouveau responsable (si 'Nouveau...' sélectionné)")
    
    qui = qui_nouveau if qui_existant == "Nouveau..." and qui_nouveau else (qui_existant if qui_existant != "Nouveau..." else "Non défini")
    quoi = col2.text_area("Remarque / Action requise (Dictée vocale recommandée)")
    quand = col3.text_input("Délai", placeholder="Ex: ASAP, 12/05...")
    
    if st.form_submit_button("Ajouter ce point (S'affichera en Bleu)"):
        st.session_state.points.append({"qui": qui, "quoi": quoi, "quand": quand, "color": "#0000FF"})
        st.rerun()

# --- 3. LE CONTRÔLE : GESTION DES POINTS ---
st.write("---")
st.subheader(f"📋 3. Ordre du jour - {nom_projet}")
for i, p in enumerate(st.session_state.points):
    colA, colB = st.columns([11, 1])
    pastille = "🔵" if p['color'] == "#0000FF" else "⚫"
    colA.write(f"{pastille} **{p['qui']}** : {p['quoi']} *(Délai: {p['quand']})*")
    
    if colB.button("❌", key=f"del_{i}", help="Supprimer ce point"):
        st.session_state.points.pop(i)
        st.rerun()

# --- 4. LA FINALISATION : EXPORTS ---
st.write("---")
st.subheader("🚀 4. Clôturer la séance")

col_word, col_json = st.columns(2)

# --- BOUTON 1 : Génération du Word DEPUIS ZÉRO ---
# On crée le document en mémoire (Page blanche)
doc = Document()

# 1. Ajout du Titre Principal
titre = doc.add_heading(f"Procès-Verbal de Séance - {nom_projet}", level=0)
doc.add_paragraph(f"Date de la séance : {date_seance.strftime('%d/%m/%Y')}")
doc.add_heading("Actions et Remarques", level=1)

# 2. Création d'un tableau (X lignes, 3 colonnes)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid' # Met des bordures noires classiques

# 3. En-têtes du tableau
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Responsable'
hdr_cells[1].text = 'Action / Remarque'
hdr_cells[2].text = 'Délai'

# 4. Remplissage du tableau avec tes points
for p in st.session_state.points:
    row_cells = table.add_row().cells
    row_cells[0].text = p['qui']
    
    # Gestion spécifique de la couleur pour la colonne "Quoi"
    p_quoi = row_cells[1].paragraphs[0]
    run = p_quoi.add_run(p['quoi'])
    
    if p['color'] == "#0000FF": # Si c'est un nouveau point
        run.font.color.rgb = RGBColor(0, 0, 255) # Texte en bleu
    else: # Si c'est un ancien point
        run.font.color.rgb = RGBColor(0, 0, 0) # Texte en noir
        run.font.bold = True # On met en gras comme dans ton code d'origine
        
    row_cells[2].text = p['quand']

# 5. Sauvegarde dans un buffer
bio_word = io.BytesIO()
doc.save(bio_word)

nom_fichier_word = f"PV_{nom_projet.replace(' ', '_')}_{date_seance}.docx"
col_word.download_button(
    label="📥 1. Télécharger le PV (Word)", 
    data=bio_word.getvalue(), 
    file_name=nom_fichier_word, 
    type="primary"
)

# --- BOUTON 2 : Génération de la mémoire (JSON) ---
json_data = json.dumps(st.session_state.points, indent=4)
nom_fichier_json = f"historique_{nom_projet.replace(' ', '_')}.json"
col_json.download_button(
    label="💾 2. Sauvegarder l'historique (JSON)", 
    data=json_data, 
    file_name=nom_fichier_json
)
